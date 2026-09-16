"""Bulk image upload: an Excel manifest + a batch of files, routed through the
same upload_image() pipeline used by the single-image upload (validation,
dedup, thumbnailing, geo-metadata) — see image_service.py."""
from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from ..config import get_storage_root
from ..models import DdaVillage, DdaZone
from .image_service import IMAGE_SOURCES, IMAGE_TYPES, upload_image
from .models import ImageLibrary
from .tree_service import get_or_create_node_by_path

logger = logging.getLogger(__name__)

# General Folder replaces Folder Path (still accepted when parsing old sheets).
TEMPLATE_HEADERS = [
    "File Name",
    "Server Filename",
    "General Folder",
    "Zone",
    "Village",
    "Image Type",
    "Year",
    "Source",
    "Capture Date",
    "Manual Bounds",
]
TEMPLATE_EXAMPLE_ROW = [
    "example_before.tif",
    "central_delhi_2025_before.tif",
    "Central Delhi/Sector 5/2025",
    "Central Delhi",
    "Sector 5",
    "GeoTIFF",
    "2025",
    "satellite",
    "2025-01-15",
    "",
]
_REQUIRED_COLUMNS = ("File Name", "General Folder", "Zone", "Village", "Image Type")
_SOURCE_CHOICES = sorted(s for s in IMAGE_SOURCES if s)


def _basename(value: str) -> str:
    """Strip quotes and take the last path segment so Explorer 'Copy as path'
    pastes still match browser upload basenames."""
    v = (value or "").strip()
    if len(v) >= 2 and v[0] == v[-1] and v[0] in ('"', "'"):
        v = v[1:-1].strip()
    return Path(v.replace("\\", "/")).name


def _normalize_general_folder(value: str) -> str:
    """Relative slash-separated path under Main folder. Reject absolute paths."""
    v = (value or "").strip()
    if len(v) >= 2 and v[0] == v[-1] and v[0] in ('"', "'"):
        v = v[1:-1].strip()
    v = v.replace("\\", "/").strip()
    if not v:
        raise ValueError("General Folder is required")
    if re.match(r"^[A-Za-z]:[/\\]", v) or v.startswith("//") or v.startswith("/"):
        raise ValueError(
            "General Folder must be relative to the server Main folder "
            "(e.g. Central Delhi/Sector 5/2025), not an absolute path"
        )
    parts = [p.strip() for p in v.split("/") if p.strip() and p.strip() != "."]
    if any(p == ".." for p in parts):
        raise ValueError("General Folder cannot contain '..'")
    if not parts:
        raise ValueError("General Folder is required")
    return "/".join(parts)


def build_bulk_upload_template() -> bytes:
    """Generate the standardized .xlsx template for bulk image upload."""
    main = str(get_storage_root())
    wb = Workbook()
    ws = wb.active
    ws.title = "Images"
    ws.append(TEMPLATE_HEADERS)
    ws.append(TEMPLATE_EXAMPLE_ROW)
    widths = (28, 32, 40, 16, 16, 14, 10, 12, 14, 22)
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    dv_type = DataValidation(
        type="list", formula1='"' + ",".join(sorted(IMAGE_TYPES)) + '"', allow_blank=False
    )
    ws.add_data_validation(dv_type)
    dv_type.add("F2:F1000")

    if _SOURCE_CHOICES:
        dv_src = DataValidation(
            type="list", formula1='"' + ",".join(_SOURCE_CHOICES) + '"', allow_blank=True
        )
        ws.add_data_validation(dv_src)
        dv_src.add("H2:H1000")

    instructions = wb.create_sheet("Instructions")
    instructions.append(["Column", "Required", "Notes"])
    instructions.append([
        "Main folder (this server)",
        "—",
        f"Images are stored under: {main}. Set STORAGE_ROOT or LOCAL_LIBRARY_ROOT on each environment so relative General Folder paths resolve correctly.",
    ])
    instructions.append([
        "File Name",
        "Yes",
        "Must match one of the files you upload with this sheet (basename). Pasting a full Explorer path is fine — only the filename is used for matching.",
    ])
    instructions.append([
        "Server Filename",
        "No",
        "Name stored on the server. Blank = File Name basename.",
    ])
    instructions.append([
        "General Folder",
        "Yes",
        "Relative path under Main folder / library tree, e.g. Central Delhi/Sector 5/2025. Absolute paths (C:\\… or /…) are rejected. Missing folders are created automatically.",
    ])
    instructions.append(["Zone", "Yes", "Must match an existing Zone name exactly (see Zone / District master)."])
    instructions.append(["Village", "Yes", "Must match an existing Village name within that Zone."])
    instructions.append(["Image Type", "Yes", "One of: " + ", ".join(sorted(IMAGE_TYPES))])
    instructions.append(["Year", "No", "Four-digit year (1900–2100). Blank may be inferred from a year segment in General Folder."])
    instructions.append([
        "Source",
        "No",
        "One of: " + (", ".join(_SOURCE_CHOICES) if _SOURCE_CHOICES else "satellite or drone"),
    ])
    instructions.append(["Capture Date", "No", "YYYY-MM-DD"])
    instructions.append(["Manual Bounds", "No", "West,South,East,North — only if the image has no embedded georeferencing."])
    instructions.column_dimensions["A"].width = 28
    instructions.column_dimensions["B"].width = 10
    instructions.column_dimensions["C"].width = 100

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@dataclass
class ManifestRow:
    row_number: int
    file_name: str
    folder_path: str
    zone: str
    village: str
    image_type: str
    capture_date: Optional[str]
    manual_bounds: Optional[str]
    server_filename: Optional[str] = None
    year: Optional[str] = None
    source: Optional[str] = None


@dataclass
class BulkDefaults:
    """Optional modal defaults applied only where Excel cells are blank."""
    year: Optional[str] = None
    source: Optional[str] = None
    image_type: Optional[str] = None
    zone: Optional[str] = None
    village: Optional[str] = None
    filenames: Dict[str, str] = field(default_factory=dict)


def parse_bulk_manifest(xlsx_bytes: bytes) -> tuple[List[ManifestRow], List[dict]]:
    """Parse the uploaded manifest. Returns (valid_rows, row_errors)."""
    try:
        wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read the Excel file. Please use the provided template.")
    ws = wb["Images"] if "Images" in wb.sheetnames else wb.worksheets[0]

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(c or "").strip() for c in next(rows_iter)]
    except StopIteration:
        raise HTTPException(status_code=400, detail="The Excel sheet is empty.")

    idx: Dict[str, Optional[int]] = {}
    for name in TEMPLATE_HEADERS:
        idx[name] = header.index(name) if name in header else None
    if idx.get("General Folder") is None and "Folder Path" in header:
        idx["General Folder"] = header.index("Folder Path")

    missing = [name for name in _REQUIRED_COLUMNS if idx.get(name) is None]
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Missing required column(s): {', '.join(missing)}. Please use the provided template.",
        )

    valid: List[ManifestRow] = []
    errors: List[dict] = []
    for row_number, raw in enumerate(rows_iter, start=2):
        if raw is None or all(c is None or str(c).strip() == "" for c in raw):
            continue

        def cell(name: str, _raw=raw) -> str:
            i = idx.get(name)
            if i is None or i >= len(_raw) or _raw[i] is None:
                return ""
            return str(_raw[i]).strip()

        file_name = _basename(cell("File Name"))
        server_filename = cell("Server Filename") or None
        if server_filename:
            server_filename = _basename(server_filename)
        zone = cell("Zone")
        village = cell("Village")
        image_type = cell("Image Type") or ""
        year = cell("Year") or None
        if year and year.endswith(".0"):
            year = year[:-2]
        source = (cell("Source") or "").strip().lower() or None
        capture_date = cell("Capture Date") or None
        manual_bounds = cell("Manual Bounds") or None

        try:
            folder_path = _normalize_general_folder(cell("General Folder"))
        except ValueError as exc:
            errors.append({"row": row_number, "file": file_name, "reason": str(exc)})
            continue

        if not file_name:
            errors.append({"row": row_number, "file": file_name, "reason": "File Name is required"})
            continue
        # Zone/Village/Image Type may be filled from modal defaults later.
        if image_type and image_type not in IMAGE_TYPES:
            errors.append({
                "row": row_number,
                "file": file_name,
                "reason": f"Image Type must be one of: {', '.join(sorted(IMAGE_TYPES))}",
            })
            continue
        if source and source not in IMAGE_SOURCES:
            errors.append({
                "row": row_number,
                "file": file_name,
                "reason": "Source must be satellite or drone",
            })
            continue
        valid.append(ManifestRow(
            row_number=row_number,
            file_name=file_name,
            folder_path=folder_path,
            zone=zone,
            village=village,
            image_type=image_type,
            capture_date=capture_date,
            manual_bounds=manual_bounds,
            server_filename=server_filename,
            year=year,
            source=source,
        ))
    return valid, errors


def _resolve_zone_village_by_name(db: Session, zone_name: str, village_name: str) -> tuple[int, int]:
    zone = db.query(DdaZone).filter(DdaZone.name.ilike(zone_name.strip())).first()
    if not zone:
        raise HTTPException(status_code=400, detail=f"Zone '{zone_name}' not found. Check the Zone / District master.")
    village = (
        db.query(DdaVillage)
        .filter(DdaVillage.zone_id == zone.id, DdaVillage.name.ilike(village_name.strip()))
        .first()
    )
    if not village:
        raise HTTPException(
            status_code=400,
            detail=f"Village '{village_name}' not found under zone '{zone_name}'. Check the Village / Location master.",
        )
    return zone.id, village.id


def zone_village_names(db: Session, zone_id: Optional[int], village_id: Optional[int]) -> tuple[str, str]:
    zone_name, village_name = "", ""
    if zone_id:
        z = db.query(DdaZone).filter(DdaZone.id == zone_id).first()
        zone_name = z.name if z else ""
    if village_id:
        v = db.query(DdaVillage).filter(DdaVillage.id == village_id).first()
        village_name = v.name if v else ""
    return zone_name, village_name


def apply_bulk_defaults(rows: List[ManifestRow], defaults: BulkDefaults) -> None:
    """Fill blank Excel fields from modal defaults (in place)."""
    for row in rows:
        key = row.file_name.strip().lower()
        if not row.server_filename and key in defaults.filenames:
            row.server_filename = defaults.filenames[key]
        if not row.year and defaults.year:
            row.year = defaults.year
        if not row.source and defaults.source:
            row.source = defaults.source
        if not row.image_type and defaults.image_type:
            row.image_type = defaults.image_type
        if not row.zone and defaults.zone:
            row.zone = defaults.zone
        if not row.village and defaults.village:
            row.village = defaults.village
        if not row.image_type:
            row.image_type = "GeoTIFF"


@dataclass
class BulkUploadResult:
    created: List[dict] = field(default_factory=list)
    duplicates: List[dict] = field(default_factory=list)
    errors: List[dict] = field(default_factory=list)
    unmatched_files: List[dict] = field(default_factory=list)


async def bulk_upload_images(
    db: Session,
    rows: List[ManifestRow],
    files: List[UploadFile],
    *,
    uploaded_by: str,
    defaults: Optional[BulkDefaults] = None,
) -> BulkUploadResult:
    result = BulkUploadResult()
    apply_bulk_defaults(rows, defaults or BulkDefaults())

    files_by_name = {(f.filename or "").strip().lower(): f for f in files}
    matched_names = set()

    for row in rows:
        key = row.file_name.strip().lower()
        file = files_by_name.get(key)
        if file is None:
            result.errors.append({
                "row": row.row_number,
                "file": row.file_name,
                "reason": "No uploaded file matches this File Name",
            })
            continue
        matched_names.add(key)
        stored_name = row.server_filename or row.file_name
        if not row.zone or not row.village:
            result.errors.append({
                "row": row.row_number,
                "file": row.file_name,
                "reason": "Zone and Village are required (set in Excel or modal defaults)",
            })
            continue
        if row.image_type not in IMAGE_TYPES:
            result.errors.append({
                "row": row.row_number,
                "file": row.file_name,
                "reason": f"Image Type must be one of: {', '.join(sorted(IMAGE_TYPES))}",
            })
            continue
        try:
            zone_id, village_id = _resolve_zone_village_by_name(db, row.zone, row.village)
            node = get_or_create_node_by_path(db, row.folder_path, created_by=uploaded_by)
            already_existed = (
                db.query(ImageLibrary)
                .filter(
                    ImageLibrary.node_id == node.id,
                    ImageLibrary.image_name == Path(stored_name).name,
                )
                .count() > 0
            )
            img = await upload_image(
                db, node.id, file,
                image_type=row.image_type,
                capture_date=row.capture_date,
                uploaded_by=uploaded_by,
                manual_bounds=row.manual_bounds,
                zone_id=zone_id,
                village_id=village_id,
                year=row.year,
                source=row.source,
                server_filename=stored_name,
            )
            entry = {
                "row": row.row_number,
                "file": row.file_name,
                "serverFilename": stored_name,
                "folder": node.node_path,
                "imageId": img.id,
            }
            (result.duplicates if already_existed else result.created).append(entry)
        except HTTPException as exc:
            db.rollback()
            result.errors.append({"row": row.row_number, "file": row.file_name, "reason": exc.detail})
        except Exception:
            db.rollback()
            logger.exception("Bulk upload failed for row %s (%s)", row.row_number, row.file_name)
            result.errors.append({
                "row": row.row_number,
                "file": row.file_name,
                "reason": "Unexpected error while uploading this file",
            })

    for name, file in files_by_name.items():
        if name not in matched_names:
            result.unmatched_files.append({"file": file.filename})

    return result
